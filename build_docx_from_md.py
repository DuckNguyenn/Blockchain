from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "docs" / "de_tai_2_esp32_sieu_am.md"
OUTPUT = ROOT / "docs" / "de_tai_2_esp32_sieu_am.docx"
BACKTICK = chr(96)
FENCE = BACKTICK * 3


def set_run_font(run, name="Calibri", size=10.5, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    table._tbl.tblPr.append(borders)


def set_repeat_header(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def add_formatted_text(paragraph, text, size=10.5):
    parts = []
    cursor = 0
    while cursor < len(text):
        start = text.find("**", cursor)
        if start == -1:
            parts.append((text[cursor:], False))
            break
        if start > cursor:
            parts.append((text[cursor:start], False))
        end = text.find("**", start + 2)
        if end == -1:
            parts.append((text[start:], False))
            break
        parts.append((text[start + 2:end], True))
        cursor = end + 2
    if not parts:
        parts = [(text, False)]
    for value, bold in parts:
        if not value:
            continue
        run = paragraph.add_run(value)
        set_run_font(run, size=size, bold=bold)


def style_table(table, widths):
    set_cell_borders(table)
    table.autofit = False
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.1
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
    header = table.rows[0]
    set_repeat_header(header)
    for cell in header.cells:
        set_cell_shading(cell, "1F4E78")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, size=9.5, bold=True, color=RGBColor(255, 255, 255))


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    style_table(table, widths)
    for row in table.rows[1:]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=9.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)
    spacer.paragraph_format.space_before = Pt(0)


def add_code_block(doc, lines):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.left_indent = Cm(0.4)
    paragraph.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(run, name="Consolas", size=8.5)
        if index < len(lines) - 1:
            run.add_break()


def parse_table_block(lines):
    rows = []
    for line in lines:
        if not line.strip().startswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(set(cell) <= {"-", ":", " ", ""} for cell in cells):
            continue
        rows.append(cells)
    return rows[0], rows[1:]


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.18

for style_name, size, before in (("Title", 20, 0), ("Heading 1", 14, 14), ("Heading 2", 12, 10)):
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(6)

title = doc.add_paragraph("Đề tài 2 - HRC Safety Log dùng ESP32 và cảm biến siêu âm", style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(14)

content = SOURCE.read_text(encoding="utf-8").splitlines()
index = 1
while index < len(content):
    stripped = content[index].strip()
    if not stripped:
        index += 1
        continue
    if stripped.startswith(FENCE):
        index += 1
        block = []
        while index < len(content) and not content[index].strip().startswith(FENCE):
            block.append(content[index])
            index += 1
        index += 1
        add_code_block(doc, block)
        continue
    if stripped.startswith("|"):
        block = []
        while index < len(content) and content[index].strip().startswith("|"):
            block.append(content[index])
            index += 1
        headers, rows = parse_table_block(block)
        if headers and "Linh kiện" in headers[0]:
            widths = [6.0, 1.8, 8.5]
        elif headers and "Thiết bị/chân" in headers[0]:
            widths = [4.5, 11.8]
        elif headers and headers[0] == "Ngày":
            widths = [1.5, 8.0, 6.8]
        else:
            widths = [16.3 / len(headers)] * len(headers)
        add_table(doc, headers, rows, widths)
        continue
    if stripped.startswith("# "):
        index += 1
        continue
    if stripped.startswith("## "):
        doc.add_heading(stripped[3:], level=1)
    elif stripped.startswith("### "):
        doc.add_heading(stripped[4:], level=2)
    elif stripped.startswith("> "):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.4)
        add_formatted_text(paragraph, stripped[2:])
    elif stripped.startswith("- "):
        paragraph = doc.add_paragraph(style="List Bullet")
        add_formatted_text(paragraph, stripped[2:])
    else:
        paragraph = doc.add_paragraph()
        add_formatted_text(paragraph, stripped)
    index += 1

doc.save(OUTPUT)
print(OUTPUT)
